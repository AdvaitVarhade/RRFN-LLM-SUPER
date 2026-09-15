import collections.abc
from pptx import Presentation

try:
    prs = Presentation('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026.pptx')
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            if hasattr(shape, "text"):
                print(f"Shape {j}: {shape.text.strip()}")
except Exception as e:
    print("Error:", e)
