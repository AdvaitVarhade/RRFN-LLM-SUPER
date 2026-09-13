import collections.abc
from pptx import Presentation

try:
    prs = Presentation('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026.pptx')
    slide_5 = prs.slides[4]
    for j, shape in enumerate(slide_5.shapes):
        print(f"Shape {j} type: {shape.shape_type}")
        if shape.has_table:
            for row in shape.table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                print("Table Row:", row_data)
except Exception as e:
    print("Error:", e)
