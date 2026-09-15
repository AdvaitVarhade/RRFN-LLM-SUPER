import collections.abc
from pptx import Presentation

try:
    prs = Presentation('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026.pptx')

    # Slide 1: Title
    prs.slides[0].shapes[1].text = "Privacy-Preserving Federated SUPER with LLM Semantic Integration"
    prs.slides[0].shapes[2].text = "Advait Varhade"
    prs.slides[0].shapes[4].text = "Project Guide"

    # Slide 3: Field of Work and Objectives
    prs.slides[2].shapes[2].text = (
        "Field of Work:\nPrivacy-Preserving Recommender Systems, Federated Learning, and Fairness.\n\n"
        "Objectives:\n"
        "1. Adapt the SUPER framework to a federated setting (FedSUPER) to protect user data.\n"
        "2. Integrate LLM-based user profiling to recover accuracy drops caused by federated sparsity.\n"
        "3. Ensure original SUPER calibration guarantees (Rmse-PC) are strictly maintained."
    )

    # Slide 4: Introduction
    prs.slides[3].shapes[2].text = (
        "- Recommender systems often suffer from popularity bias, overshadowing long-tail items.\n"
        "- The SUPER framework effectively mitigates this by merging head and tail recommendations via user-specific blueprints.\n"
        "- However, SUPER is centralized, creating privacy risks. Federated Learning solves privacy but degrades recommendation accuracy.\n"
        "- This project integrates LLM semantic profiles into FedSUPER to achieve both privacy and high accuracy."
    )

    # Slide 5: Literature Survey
    table = prs.slides[4].shapes[6].table
    # Row 1
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "SUPER: Smart User-centric Popularity Exposure Reduction..."
    table.cell(1, 2).text = "IEEE Access"
    table.cell(1, 3).text = "2026"
    table.cell(1, 4).text = "Introduced blueprint merging for popularity bias; but it is centralized."
    # Row 2
    table.cell(2, 0).text = "2"
    table.cell(2, 1).text = "Federated Neural Collaborative Filtering"
    table.cell(2, 2).text = "Knowl.-Based Syst."
    table.cell(2, 3).text = "2021"
    table.cell(2, 4).text = "Preserves privacy via local user embeddings; suffers from data sparsity."

    # Slide 6: Outcome of Literature Survey
    prs.slides[5].shapes[2].text = (
        "Research Gaps Identified:\n\n"
        "- No existing work provides a mathematically guaranteed popularity calibration (like SUPER) in a fully federated, privacy-preserving manner.\n"
        "- Traditional federated learning causes severe accuracy drops on long-tail items due to client-side data sparsity.\n"
        "- LLM semantic integration has not been explored as a mechanism to rescue accuracy specifically within the constraints of strict popularity-calibration blueprints."
    )

    # Slide 7: Problem Definition
    prs.slides[6].shapes[2].text = (
        "Objectives:\n"
        "- Implement a Federated SUPER (FedSUPER) baseline.\n"
        "- Integrate LLM profile semantics (Sentence-BERT) into the federated pipeline.\n"
        "- Evaluate models on the ML-1M dataset.\n\n"
        "Scope of the Study:\n"
        "- Focuses on the MovieLens-1M dataset to evaluate implicit feedback recommendations.\n"
        "- Primary evaluation metrics: Recall@20 (for accuracy) and Rmse-PC (to validate calibration preservation)."
    )

    # Slide 8: Methodology/Work Plan
    prs.slides[7].shapes[2].text = (
        "1. Baseline Development: Implement centralized BPR-MF and SUPER.\n"
        "2. Federated Core (FedNCF): Distribute training with shared item embeddings and private user embeddings.\n"
        "3. FedSUPER Integration: Apply SUPER blueprint logic to federated client outputs.\n"
        "4. LLM Semantic Profiling: Generate item and user profiles using Sentence-BERT locally.\n"
        "5. Fusion & Merging: Blend LLM scores with FedNCF scores before the blueprint merge.\n"
        "6. Final Evaluation: Compute Recall, Rmse-PC, MRMC, and Fairness metrics."
    )

    # Slide 9: Design Standards & Realistic Constraints
    prs.slides[8].shapes[2].text = (
        "Design Standards:\n"
        "- Object-oriented, modular pipeline architecture in PyTorch.\n"
        "- Reproducible evaluation standards based on Yavru et al.\n\n"
        "Realistic Constraints:\n"
        "- Federated learning entails severe data sparsity on the client side, complicating long-tail learning.\n"
        "- Compute overhead for generating LLM embeddings locally on resource-constrained user devices (mitigated via lightweight MiniLM models)."
    )

    # Slide 10: Project Demonstration
    prs.slides[9].shapes[2].text = (
        "Work Done:\n"
        "- Centralized baselines (BPR-MF, SUPER) successfully implemented and validated.\n"
        "- FedSUPER baseline evaluated. Proved that blueprint merges work perfectly in federated settings, maintaining perfect calibration (Rmse-PC = 0.055).\n"
        "- LLM semantic profiling integrated into FedSUPER.\n"
        "- Results: LLM-FedSUPER recovered 91% of the federated recall gap (Recall=0.037 vs centralized 0.038) while maintaining identical popularity calibration.\n"
        "- Proven that SUPER's calibration is a property of the algorithm, invariant to the federated backbone."
    )

    # Slide 11: Work to be done
    prs.slides[10].shapes[2].text = (
        "Work to be done:\n"
        "- Formal documentation of the privacy-utility trade-offs and composite FDN-A metrics.\n"
        "- Finalizing the project report and drafting the research paper.\n"
        "- Code cleanup, documentation, and preparation for open-source publication."
    )

    # Slide 12: Time Line
    prs.slides[11].shapes[2].text = (
        "- Phase 1: Literature review, dataset preparation, and environment setup.\n"
        "- Phase 2: Centralized baseline implementation and FedNCF core.\n"
        "- Phase 3: FedSUPER integration and LLM profiling.\n"
        "- Phase 4: Final metric evaluation, reporting, and review preparation."
    )

    # Slide 13: References
    prs.slides[12].shapes[2].text = (
        "1. Yavru, S. et al. (2026). SUPER: Smart User-centric Popularity Exposure Reduction for Fair and Diverse Recommendations, IEEE Access.\n"
        "2. Perifanis, V. et al. (2021). Federated Neural Collaborative Filtering, Knowledge-Based Systems.\n"
        "3. Reimers, N. et al. (2019). Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks, EMNLP."
    )

    prs.save('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026_Updated.pptx')
    print("Successfully updated PPT.")
except Exception as e:
    print("Error:", e)
