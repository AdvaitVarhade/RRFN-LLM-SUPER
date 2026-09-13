import collections.abc
from pptx import Presentation

try:
    prs = Presentation('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026.pptx')
    slide_5 = prs.slides[4]
    table_shape = slide_5.shapes[6]
    print("Table dimensions:")
    print(f"Left: {table_shape.left}")
    print(f"Top: {table_shape.top}")
    print(f"Width: {table_shape.width}")
    print(f"Height: {table_shape.height}")
    
    # Also find layout for slide 5
    print("Layout name:", slide_5.slide_layout.name)
except Exception as e:
    print("Error:", e)
