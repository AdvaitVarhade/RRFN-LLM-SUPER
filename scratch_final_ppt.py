import json
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def replace_footer(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if "BCLE497J" in run.text:
                            run.text = run.text.replace("BCLE497J", "BCSE497J")

def apply_table_style(table):
    # Just basic formatting for the header and rows
    pass

def set_cell_text(cell, text):
    cell.text = text
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(12)

def populate_table(slide, papers, start_idx):
    # Remove existing table if any (like on the original Slide 5)
    shapes_to_remove = []
    for shape in slide.shapes:
        if shape.has_table:
            shapes_to_remove.append(shape)
    for shape in shapes_to_remove:
        # In python-pptx, deleting a shape is a bit hacky via the XML element
        shape.element.getparent().remove(shape.element)

    # Dimensions
    left = 3556794
    top = 2990056
    width = 8382000
    height = 3352800
    
    table_shape = slide.shapes.add_table(rows=6, cols=5, left=left, top=top, width=width, height=height)
    table = table_shape.table
    
    headers = ['S. No.', 'Paper Title', 'Venue', 'Year', 'Important Findings']
    for col_idx, header in enumerate(headers):
        set_cell_text(table.cell(0, col_idx), header)
        
    for i in range(5):
        if i < len(papers):
            paper = papers[i]
            set_cell_text(table.cell(i+1, 0), str(start_idx + i + 1))
            set_cell_text(table.cell(i+1, 1), paper['title'])
            set_cell_text(table.cell(i+1, 2), paper['venue'])
            set_cell_text(table.cell(i+1, 3), paper['year'])
            set_cell_text(table.cell(i+1, 4), paper['abstract'])

def main():
    prs = Presentation('c:/d_drive/projects/Project1/BCLE497J Project-I_Review Presentation 1_2026.pptx')
    
    # 1. Load papers
    with open('c:/d_drive/projects/Project1/papers_15.json', 'r', encoding='utf-8') as f:
        papers = json.load(f)
        
    # 2. Setup 3 slides for Lit Survey
    # Slide 5 (index 4) is the original Lit Survey. We will add 2 new slides right after it.
    # It's easiest to just add 2 slides at the end and reorder them in the XML.
    slide_5 = prs.slides[4]
    
    # Create two new slides for the rest of the lit survey
    layout = slide_5.slide_layout
    slide_6 = prs.slides.add_slide(layout)
    slide_7 = prs.slides.add_slide(layout)
    
    # Set titles for the new slides
    slide_6.shapes.title.text = "Literature Survey (Cont.)"
    slide_7.shapes.title.text = "Literature Survey (Cont.)"
    
    # Populate the tables
    populate_table(slide_5, papers[0:5], 0)
    populate_table(slide_6, papers[5:10], 5)
    populate_table(slide_7, papers[10:15], 10)
    
    # Move the new slides to positions 5 and 6 (0-indexed 5 and 6)
    # The slides we just added are at the end: -2 and -1
    sldIdLst = prs.slides._sldIdLst
    sldIdLst.insert(5, sldIdLst[-2])
    sldIdLst.insert(6, sldIdLst[-1])

    # 3. Replace footer
    replace_footer(prs)
    
    # 4. Modify Project Demonstration (now shifted by 2, so it was index 9, now it is index 11)
    demo_slide = prs.slides[11]
    # Remove any results, only put contents mentioned in review 2 (architecture/module/methodology setup)
    demo_slide.shapes[2].text = (
        "Experimental Design & Implementation:\n"
        "- Baseline data pipeline completed for ML-1M dataset partitioning.\n"
        "- Federated Collaborative Filtering (FedNCF) core architecture setup.\n"
        "- LLM Sentence-BERT integrated for initial semantic profiling.\n"
        "- SUPER blueprint merging algorithm implemented in federated clients.\n"
        "- The experimental environment is fully configured and ready for executing large-scale evaluations."
    )
    
    # 5. Modify Work to be Done (was index 10, now index 12)
    work_slide = prs.slides[12]
    work_slide.shapes[2].text = (
        "Work to be done:\n"
        "- Execute full experimental evaluation to measure Recall@20 and Rmse-PC.\n"
        "- Finalize the comparison showing LLM-FedSUPER recovering the recall gap while maintaining calibration.\n"
        "- Document the privacy-utility trade-offs thoroughly.\n"
        "- Prepare the final report detailing the empirical findings."
    )
    
    # 6. References (was index 12, now index 14)
    ref_slide = prs.slides[14]
    ref_text = ""
    for i, p in enumerate(papers):
        ref_text += f"{i+1}. {p['title']}, {p['venue']}, {p['year']}.\n"
    ref_slide.shapes[2].text = ref_text
    # Reduce font size to fit 15 references
    for p in ref_slide.shapes[2].text_frame.paragraphs:
        p.font.size = Pt(10)
        
    prs.save('c:/d_drive/projects/Project1/BCSE497J Project-I_Review Presentation 1_2026.pptx')
    print("Done editing PPT!")

if __name__ == '__main__':
    main()
